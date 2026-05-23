
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

places_labels=["🏆 1st","🥈 2nd","🥉 3rd","4th","5th","6th","7th","8th","9th","10th","11th","12th","13th","14th"]
A_mode=[all_teams_A[np.argmax(pos_A[:,p])] for p in range(7)]
B_mode=[all_teams_B[np.argmax(pos_B[:,p])] for p in range(7)]
ml_fin=sorted(all_teams,key=lambda t:finish[t][0]+finish[t][1],reverse=True)[:4]
fin2=sorted(ml_fin[:2],key=lambda t:finish[t][0],reverse=True)
bro2=sorted(ml_fin[2:4],key=lambda t:finish[t][2],reverse=True)

# ── STANDINGS ──
def make_standings(teams, stats):
    srt=sorted(teams,key=lambda t:(stats[t]['Pts'],stats[t]['GF']-stats[t]['GA'],stats[t]['GF']),reverse=True)
    rows=[]
    for i,t in enumerate(srt):
        s=stats[t]
        rows.append({"Pos":i+1,"Team":t,"G":s['G'],"W":s['W'],"D":s['D'],"L":s['L'],
                     "GF":s['GF'],"GA":s['GA'],"GD":s['GF']-s['GA'],"Pts":s['Pts'],"Games Left":6-s['G']})
    return pd.DataFrame(rows)

df_sA=make_standings(all_teams_A,statsA)
df_sB=make_standings(all_teams_B,statsB)

# ── TEAM PARAMETERS ──
model_rows=[]
for grp,teams,att_s,def_s,ratt,rdef,alpha,avg in [
    ("A",all_teams_A,attA,defA,rawattA,rawdefA,alphaA,avgA),
    ("B",all_teams_B,attB,defB,rawattB,rawdefB,alphaB,avgB)]:
    for t in teams:
        s=statsA[t] if grp=="A" else statsB[t]; g=max(s['G'],1); a_v=g/(g+5)
        model_rows.append({"Team":t,"Group":grp,"Games":s['G'],"GF":s['GF'],"GA":s['GA'],
                           "Raw Att":round(ratt[t],3),"Raw Def":round(rdef[t],3),
                           "α (shrink)":round(a_v,3),
                           "Adj Att":round(att_s[t],3),"Adj Def":round(def_s[t],3),
                           "Quality Score":round(all_quality[t],3),
                           "Grp Avg λ":round(avg,3)})
df_model=pd.DataFrame(model_rows)

# ── GLOBAL QUALITY RANKING ──
qual_rows=[]
for rank,(t,q) in enumerate(sorted(all_quality.items(),key=lambda x:-x[1]),1):
    grp="A" if t in all_teams_A else "B"
    s=statsA[t] if grp=="A" else statsB[t]
    qual_rows.append({"Overall Rank":rank,"Team":t,"Group":grp,
                      "Quality Score":round(q,3),
                      "Adj Att":round(all_att[t],3),"Adj Def":round(all_def[t],3),
                      "Formula":"att/def × avg (norm)"})
df_qual=pd.DataFrame(qual_rows)

# ── BRIDGE RSI MATRIX (cross-group) ──
bridge_rows=[]
for h in all_teams_A:
    row={"Grp A \\ Grp B (RSI_H)":h}
    for a in all_teams_B:
        rh,ra=bridge_rsi(h,a,lkA,lkB,all_quality,all_teams_A,all_teams_B)
        row[a]=round(rh,3)
    bridge_rows.append(row)
df_bridge=pd.DataFrame(bridge_rows)

# ── KO PREDICTIONS ──
ko_rows=[]
for num,stage,h,a,when in [
    (43,"SF1 (A1 vs B2)",A_mode[0],B_mode[1],"Sun 13:50"),
    (44,"SF2 (B1 vs A2)",B_mode[0],A_mode[1],"Sun 13:50"),
    (45,"Bronze",bro2[0],bro2[1],"Sun 16:40"),
    (46,"Final",fin2[0],fin2[1],"Sun 16:40"),
    (47,"13-14th",A_mode[6],B_mode[6],"Sun 14:25"),
    (48,"11-12th",A_mode[5],B_mode[5],"Sun 14:25"),
    (49,"9-10th",A_mode[4],B_mode[4],"Sun 15:15"),
    (50,"7-8th",A_mode[3],B_mode[3],"Sun 15:15"),
    (51,"5-6th",A_mode[2],B_mode[2],"Sun 15:50")]:
    lh,la,ml_h,ml_a,ph,pd_,pa,rh,ra,rt=get_pred_full(h,a)
    tot=ph+pa; pw=ph/tot if tot>0 else 0.5; pl=pa/tot if tot>0 else 0.5
    ko_rows.append({"#":num,"Stage":stage,"When":when,"Home":h,"Away":a,
                    "λH":round(lh,2),"λA":round(la,2),"Score":f"{ml_h}–{ml_a}",
                    "P(Home Win KO)%":round(pw*100,1),"P(Draw@60min)%":round(pd_*100,1),
                    "P(Away Win KO)%":round(pl*100,1),
                    "RSI_H":round(rh,3),"RSI_A":round(ra,3),"RSI Type":rt})
df_ko=pd.DataFrame(ko_rows)

# ── GROUP POSITION PROBS ──
pos_rows=[]
for i,t in enumerate(all_teams_A):
    s=statsA[t]; row={"Team":t,"Group":"A","Pts":s['Pts'],"Quality":round(all_quality[t],2)}
    for p in range(7): row[f"P(Pos{p+1})%"]=round(pos_A[i,p]*100,1)
    pos_rows.append(row)
for i,t in enumerate(all_teams_B):
    s=statsB[t]; row={"Team":t,"Group":"B","Pts":s['Pts'],"Quality":round(all_quality[t],2)}
    for p in range(7): row[f"P(Pos{p+1})%"]=round(pos_B[i,p]*100,1)
    pos_rows.append(row)
df_pos=pd.DataFrame(pos_rows)

# ── TOURNAMENT FINISH ──
finish_rows=[]
for t in all_teams:
    grp="A" if t in all_teams_A else "B"; s=statsA[t] if grp=="A" else statsB[t]
    row={"Team":t,"Group":grp,"Pts":s['Pts'],"Quality":round(all_quality[t],2)}
    for p,lbl in enumerate(places_labels): row[lbl]=round(finish[t][p]/N*100,1)
    row["Semi%"]=round(sum(finish[t][:4])/N*100,1)
    row["Final%"]=round(sum(finish[t][:2])/N*100,1)
    row["Bronze%"]=round(finish[t][2]/N*100,1)
    row["Champion%"]=round(finish[t][0]/N*100,1)
    finish_rows.append(row)
df_finish=pd.DataFrame(finish_rows).sort_values("Champion%",ascending=False).reset_index(drop=True)

# ── FORMULA SHEET ──
formula_rows=[
    ("1","League avg","avg = Σ(GF)/Σ(G)","Grp A=4.143 | Grp B=2.643 | KO=3.393"),
    ("2","Raw Attack","raw_att = (GF/G) / avg","Goals per game normalized to league avg"),
    ("3","Raw Defence","raw_def = (GA/G) / avg","Conceded per game normalized to league avg"),
    ("4","Shrinkage α","α = G / (G + k),  k=5","4 games → α=0.44 (44% data, 56% mean prior)"),
    ("5","Adj ratings","adj_att = α×raw_att + (1−α)×1.0","Shrinks extremes caused by small sample"),
    ("6","Quality Score","Q = adj_att / adj_def × avg","Single ranking metric, normalized across groups"),
    ("7","Base λ","λH = adj_att_H × adj_def_A × avg","Poisson rate for home goals"),
    ("8","Direct RSI (same group)","RSI_H = exp(mean(log(gf_H_vs_C / gf_A_vs_C)))^β, β=0.3","Geometric mean over all common opponents C"),
    ("9","Bridge RSI (cross-group)","For each opp OH of H, find closest peer in A's group by quality score","Compare H's goals vs OH to A's goals vs peer(OH)"),
    ("10","Bridge weighting","Weight = exp(−|Q_OH − Q_peer| / Q_OH)","Penalizes dissimilar peer matchups"),
    ("11","Bridge blend","λH_final = λH_base × BridgeRSI_H^β, β=0.25","β=0.25 (more conservative than direct RSI)"),
    ("12","Poisson PMF","P(X=k) = e^(−λ) × λ^k / k!","Prob of scoring exactly k goals"),
    ("13","Win/Draw/Loss","Summed over 20×20 scoreline grid","P(H>A), P(H=A), P(H<A)"),
    ("14","KO prob","P(HomeWin KO) = P(H>A)/(P(H>A)+P(A>H))","No draws in KO; OT/shootout if tied"),
    ("15","Tiebreak 1-8","Pts → H2H pts → H2H GD → H2H GF → GD → GF → PenMin → Dice","Full official rules implemented in simulation"),
    ("16","Simulation","100,000 Monte Carlo runs","Poisson goals → tiebreak → KO bracket → finish positions"),
]
df_formulas=pd.DataFrame(formula_rows,columns=["Step","Component","Formula","Notes"])

print("All DataFrames ready. Writing Excel...")
